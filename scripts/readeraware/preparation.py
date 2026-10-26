"""Main-agent document tools. Never invoked by the five child-agent calls."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

from .inputs import load_sources
from .util import ValidationError, check_no_credentials, digest, json_write, text_write


def prepare_inputs(paths, out_dir, *, project_root=None, main_tex=None,
                   docx_engine='python-docx', pandoc_bin=None):
    out = Path(out_dir).resolve()
    if out.exists():
        raise ValidationError('Preparation directory already exists; use a new directory')
    # Extraction may create only explicitly requested resource files in this new directory.
    sources, provenance = load_sources(paths, project_root=project_root, main_tex=main_tex,
        prepared_dir=out, docx_engine=docx_engine, pandoc_bin=pandoc_bin)
    out.mkdir(parents=True, exist_ok=True)
    base = (out / 'originals').resolve()
    for record in provenance['files']:
        source = Path(record['path'])
        target = (base / record['relative_path']).resolve()
        if not target.is_relative_to(base):
            raise ValidationError('Original snapshot path escapes preparation directory')
        if digest(source) != record['sha256']:
            raise ValidationError('Source changed during preparation')
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, target)
        record['original_path'] = str(source)
        record['path'] = str(target)
        record['snapshot_path'] = target.relative_to(out).as_posix()
    material = '\n\n'.join('## ' + s['source_id'] + ': ' + s['name'] + '\n\n' + s['text'] for s in sources)
    check_no_credentials(material, 'prepared material')
    text_write(out / 'materials.md', material)
    json_write(out / 'sources.json', sources)
    json_write(out / 'provenance.json', provenance)
    checklist = {
        'status': 'requires_main_agent_verification',
        'machine_extraction_is_not_fidelity_certification': True,
        'material_sha256': digest(out / 'materials.md'),
        'editable_equations_verified': False,
        'editable_tables_verified': False,
        'figure_files_and_captions_verified': False,
        'citations_and_references_verified': False,
        'source_content_and_scope_preserved': False,
        'all_pdf_pages_visually_checked_if_applicable': False,
        'utf8_readback_passed': (out / 'materials.md').read_bytes().decode('utf-8') == material,
        'main_agent_notes': '',
    }
    json_write(out / 'preparation_checklist.json', checklist)
    return {'status': checklist['status'], 'material': str(out / 'materials.md'),
        'provenance': str(out / 'provenance.json'),
        'checklist': str(out / 'preparation_checklist.json'),
        'findings': provenance.get('preparation_findings', []),
        'warnings': provenance.get('warnings', []), 'model_calls': 0}


def inspect_delivery(paths):
    """Reopen files and check encoding/container integrity; visual QA stays explicit."""
    results = []
    for raw in paths:
        path = Path(raw).resolve()
        if not path.is_file():
            raise ValidationError('Deliverable not found: ' + path.name)
        record = {'name': path.name, 'sha256': digest(path), 'bytes': path.stat().st_size}
        suffix = path.suffix.lower()
        if suffix in ('.md', '.txt', '.tex', '.bib', '.json', '.csv'):
            text = path.read_bytes().decode('utf-8-sig', errors='strict')
            record.update(utf8_readback=True, cjk_characters=_cjk_count(text))
        elif suffix == '.docx':
            import zipfile
            from docx import Document
            from xml.etree import ElementTree as ET
            with zipfile.ZipFile(path) as package:
                if package.testzip():
                    raise ValidationError('Corrupt DOCX container')
                for name in package.namelist():
                    if name.endswith('.xml'):
                        ET.fromstring(package.read(name))
                xml = ET.fromstring(package.read('word/document.xml'))
                text = ''.join(xml.itertext())
            doc = Document(path)  # Actual parser reopen, not filename-only existence.
            record.update(container_reopen=True, xml_parse=True,
                paragraphs=len(doc.paragraphs), tables=len(doc.tables),
                editable_math_nodes=sum(1 for e in xml.iter() if e.tag.endswith('}oMath')),
                cjk_characters=_cjk_count(text), target_application_open='pending_main_agent')
        elif suffix == '.pdf':
            from pypdf import PdfReader
            pdf = PdfReader(path, strict=True)
            if pdf.is_encrypted or not pdf.pages:
                raise ValidationError('PDF is encrypted or has no pages')
            text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
            record.update(container_reopen=True, pages=len(pdf.pages),
                cjk_characters=_cjk_count(text), page_render_and_visual_qa='pending_main_agent')
        else:
            raise ValidationError('Unsupported verification type: ' + suffix)
        if '\ufffd' in text or '\x00' in text:
            raise ValidationError('Replacement character or NUL found in ' + path.name)
        record['readback_passed'] = True
        results.append(record)
    return {'status': 'readback_passed', 'files': results, 'model_calls': 0,
        'delivery_complete': False,
        'remaining_main_agent_checks': ['Actual target-application open where available',
            'TeX compilation and reference resolution when LaTeX is delivered',
            'Render every PDF/Word page and inspect CJK glyphs, formulas, tables, figures and layout',
            'Compare deliverable content and required structures with frozen Markdown']}


def _cjk_count(text):
    return sum('\u3400' <= c <= '\u9fff' or '\U00020000' <= c <= '\U0003134f' for c in text)
