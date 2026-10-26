"""Synthetic input-adapter checks; no model calls or native converters."""
from __future__ import annotations

import hashlib
import importlib.util
import io
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zlib

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))

from readeraware.inputs import InputError, load_sources

HAS_PDF = bool(importlib.util.find_spec('pypdf'))
HAS_DOCX = bool(importlib.util.find_spec('docx'))
HAS_REPORTLAB = bool(importlib.util.find_spec('reportlab'))


def tiny_png():
    """A synthetic 2x2 RGB fixture; no external files or image services."""
    def chunk(kind, data):
        return struct.pack('!I', len(data)) + kind + data + struct.pack('!I', zlib.crc32(kind + data) & 0xffffffff)
    rows = b'\x00\xff\x00\x00\x00\xff\x00' * 2
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('!IIBBBBB', 2, 2, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(rows)) + chunk(b'IEND', b'')


class AdapterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=Path(__file__).parent)
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def write(self, name, text):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(text.encode('utf-8'))
        return path

    def test_plain_text_preserves_bom_newlines_unicode_and_ids(self):
        first = self.write('a.md', '\ufeff标题\r\nFirst.\r\nSecond.\n')
        second = self.write('b.txt', 'Other.\r\n')
        sources, provenance = load_sources([first, second])
        self.assertEqual(sources, [{'source_id': 'S001', 'name': 'a.md',
            'text': '标题\r\nFirst.\r\nSecond.\n'}, {'source_id': 'S002',
            'name': 'b.txt', 'text': 'Other.\r\n'}])
        for source, record in zip((first, second), provenance['files']):
            self.assertEqual(record['sha256'], hashlib.sha256(source.read_bytes()).hexdigest())
            self.assertEqual(record['path'], str(source.resolve()))
            self.assertEqual(record['relative_path'], record['source_id'] + '/' + source.name)
        self.assertTrue(provenance['sources'][0]['extraction']['bom_stripped'])
        self.assertEqual(set(sources[0]), {'source_id', 'name', 'text'})
        self.assertFalse(provenance['preparation']['editable_markdown_verified'])

    def test_distinct_folders_do_not_collide_in_snapshot_paths(self):
        first = self.write('a/draft.md', 'First')
        second = self.write('b/draft.md', 'Second')
        _, provenance = load_sources([first, second])
        self.assertEqual([r['relative_path'] for r in provenance['files']],
                         ['S001/draft.md', 'S002/draft.md'])

    def test_tex_nested_includes_import_assets_and_style_dependencies(self):
        main = self.write('project/main.tex', r'''\documentclass{local}
\usepackage{graphicx}
\graphicspath{{figures/}}
Before.
\input{sections/first}
\includegraphics{chart}
\bibliography{references,missing}
\bibliographystyle{plain}
After.''')
        self.write('project/sections/first.tex', 'First.\\input{second}\n')
        self.write('project/sections/second.tex', 'Second.\\import{../other/}{third.tex}')
        self.write('project/other/third.tex', 'Third.')
        self.write('project/local.cls', '\\RequirePackage{mystyle}\n')
        self.write('project/mystyle.sty', '\\input{helper.tex}')
        self.write('project/helper.tex', '% style helper\n')
        self.write('project/references.bib', '@article{x,title={Reference}}')
        graphic = self.root / 'project/figures/chart.png'
        graphic.parent.mkdir(parents=True)
        graphic.write_bytes(b'synthetic image dependency')
        sources, provenance = load_sources([main])
        self.assertIn('Before.\nFirst.Second.Third.', sources[0]['text'])
        self.assertIn('After.', sources[0]['text'])
        self.assertNotIn('\\input{sections/first}', sources[0]['text'])
        names = {Path(item['path']).name: item['role'] for item in provenance['files']}
        self.assertEqual(names['first.tex'], 'tex-include')
        self.assertEqual(names['third.tex'], 'tex-include')
        self.assertEqual(names['helper.tex'], 'tex-include')
        self.assertEqual(names['local.cls'], 'tex-style')
        self.assertEqual(names['mystyle.sty'], 'tex-style')
        self.assertEqual(names['references.bib'], 'tex-bibliography')
        self.assertEqual(names['chart.png'], 'tex-graphic')
        self.assertTrue(any('missing' in notice for notice in provenance['warnings']))
        references = provenance['sources'][0]['extraction']['references']
        self.assertEqual(next(r for r in references if r['requested'] == 'graphicx')['status'], 'system-package')
        self.assertFalse(any('graphicx' in notice for notice in provenance['warnings']))
        self.assertTrue(all(not Path(record['relative_path']).is_absolute() for record in provenance['files']))

    def test_comments_verbatim_and_escaped_backslashes_do_not_load_files(self):
        main = self.write('main.tex', r'''\documentclass{article}
% \input{comment-missing}
\verb|\input{inline-missing}| \lstinline[language=TeX]|\include{inline2-missing}|
\begin{verbatim}
\input{verbatim-missing}
\end{verbatim}
\begin{minted}{tex}
\input{minted-missing}
\end{minted}
\\input{escaped-missing}
Escaped percent \% is text. \input{yes}
''')
        self.write('yes.tex', 'Included correctly.')
        sources, provenance = load_sources([main])
        self.assertIn('Included correctly.', sources[0]['text'])
        self.assertEqual({Path(r['path']).name for r in provenance['files']}, {'main.tex', 'yes.tex'})

    def test_include_cycles_fail_with_chain(self):
        main = self.write('main.tex', '\\input{one}')
        self.write('one.tex', '\\include{main}')
        with self.assertRaisesRegex(InputError, r'cycle: main.tex -> one.tex -> main.tex'):
            load_sources([main])

    def test_supplied_bbl_is_retained_only_for_active_bibliography(self):
        main = self.write('project/paper.tex', '\\documentclass{article}\n\\input{ending}')
        self.write('project/ending.tex', '\\printbibliography\n')
        supplied_bbl = self.write('project/paper.bbl', '% supplied bibliography\r\n\\bibitem{x} Entry.\r\n')
        self.write('project/unrelated.bbl', 'Do not copy unrelated build artifacts.')
        _, provenance = load_sources([main])
        bbl_files = [item for item in provenance['files'] if item['role'] == 'tex-bibliography-output']
        self.assertEqual(len(bbl_files), 1)
        self.assertEqual(bbl_files[0]['path'], str(supplied_bbl.resolve()))
        self.assertEqual(bbl_files[0]['relative_path'], 'S001/paper.bbl')
        self.assertEqual(bbl_files[0]['sha256'], hashlib.sha256(supplied_bbl.read_bytes()).hexdigest())
        self.write('project/ending.tex', '% \\printbibliography\nNo active bibliography.')
        _, provenance = load_sources([main])
        self.assertFalse(any(item['role'] == 'tex-bibliography-output' for item in provenance['files']))

    def test_missing_required_include_is_not_silently_omitted(self):
        main = self.write('main.tex', '\\input{missing}')
        with self.assertRaisesRegex(InputError, 'Missing local TeX dependency'):
            load_sources([main])

    def test_paths_outside_project_are_rejected_even_when_existing(self):
        self.write('outside.tex', 'Private material')
        main = self.write('project/main.tex', '\\input{../outside}')
        with self.assertRaisesRegex(InputError, 'escapes project root'):
            load_sources([main])
        graphic = self.write('project/graphic.tex', '\\includegraphics{../outside.tex}')
        with self.assertRaisesRegex(InputError, 'escapes project root'):
            load_sources([graphic])

    def test_explicit_project_root_allows_shared_files_inside_boundary(self):
        main = self.write('project/chapter/main.tex', '\\input{../shared}')
        self.write('project/shared.tex', 'Shared project content.')
        sources, provenance = load_sources([main], project_root=self.root / 'project')
        self.assertEqual(sources[0]['text'], 'Shared project content.')
        self.assertEqual(provenance['files'][0]['relative_path'], 'S001/chapter/main.tex')

    def test_remote_and_dynamic_includes_are_not_fetched_or_guessed(self):
        remote = self.write('remote.tex', '\\input{https://example.invalid/source}')
        with self.assertRaisesRegex(InputError, 'Only project-local'):
            load_sources([remote])
        dynamic = self.write('dynamic.tex', '\\input{\\chaptername}')
        with self.assertRaisesRegex(InputError, 'nonliteral'):
            load_sources([dynamic])

    def test_main_discovery_selection_and_ambiguity(self):
        self.write('project/first.tex', '\\documentclass{article}\nFirst main.')
        self.write('project/chapter.tex', 'Chapter only.')
        sources, _ = load_sources([self.root / 'project'])
        self.assertEqual(sources[0]['name'], 'first.tex')
        self.write('project/second.tex', '\\documentclass{article}\nSecond main.')
        with self.assertRaisesRegex(InputError, 'explicit main_tex'):
            load_sources([self.root / 'project'])
        sources, _ = load_sources([self.root / 'project'], main_tex='second.tex')
        self.assertEqual(sources[0]['name'], 'second.tex')

    def test_main_override_with_multiple_tex_sources_is_rejected(self):
        first = self.write('a.tex', 'First')
        second = self.write('b.tex', 'Second')
        with self.assertRaisesRegex(InputError, 'exactly one TeX'):
            load_sources([first, second], main_tex='a.tex')

    def test_empty_unknown_missing_and_legacy_inputs_fail_clearly(self):
        for name, text, error in [('empty.txt', ' \r\n', 'no extractable'),
                                  ('legacy.doc', 'legacy', 'convert.*docx'),
                                  ('unsupported.zip', 'zip', 'Unsupported')]:
            path = self.write(name, text)
            with self.subTest(name=name), self.assertRaisesRegex(InputError, error):
                load_sources([path])
        with self.assertRaisesRegex(InputError, 'at least one'):
            load_sources([])
        with self.assertRaisesRegex(InputError, 'does not exist'):
            load_sources([self.root / 'nonexistent.md'])

    def test_missing_optional_dependencies_have_actionable_errors(self):
        pdf = self.write('sample.pdf', '%PDF-test')
        docx = self.write('sample.docx', 'fake docx')
        with patch.dict(sys.modules, {'pypdf': None}):
            with self.assertRaisesRegex(InputError, 'optional dependency pypdf'):
                load_sources([pdf])
        with patch.dict(sys.modules, {'docx': None}):
            with self.assertRaisesRegex(InputError, 'optional dependency python-docx'):
                load_sources([docx])

    @unittest.skipUnless(HAS_PDF and HAS_REPORTLAB, 'pypdf and reportlab are optional')
    def test_pdf_page_order_metadata_and_partial_blank_warning(self):
        from reportlab.pdfgen.canvas import Canvas
        path = self.root / 'ordered.pdf'
        canvas = Canvas(str(path))
        canvas.setTitle('Synthetic extraction test')
        canvas.drawString(40, 750, 'First page finding.')
        canvas.showPage()
        canvas.showPage()
        canvas.drawString(40, 750, 'Third page finding.')
        canvas.save()
        sources, provenance = load_sources([path])
        text = sources[0]['text']
        self.assertLess(text.index('First page'), text.index('Third page'))
        self.assertIn('[PDF page 2]', text)
        extraction = provenance['sources'][0]['extraction']
        self.assertEqual(extraction['page_count'], 3)
        self.assertEqual(extraction['empty_text_pages'], [2])
        self.assertEqual(extraction['metadata']['/Title'], 'Synthetic extraction test')
        self.assertTrue(any('without extractable text: 2' in n for n in provenance['warnings']))
        finding = next(item for item in provenance['preparation_findings'] if item['code'] == 'pdf_editable_fidelity_unknown')
        self.assertEqual(finding['formula_presence'], 'unknown')
        self.assertEqual(finding['table_presence'], 'unknown')
        self.assertEqual(finding['status'], 'unresolved')

    @unittest.skipUnless(HAS_PDF, 'pypdf is optional')
    def test_blank_and_encrypted_pdf_fail_clearly(self):
        from pypdf import PdfWriter
        blank = self.root / 'blank.pdf'
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        writer.write(str(blank))
        with self.assertRaisesRegex(InputError, 'no extractable text'):
            load_sources([blank])
        encrypted = self.root / 'encrypted.pdf'
        writer.encrypt('nonempty-password')
        writer.write(str(encrypted))
        with self.assertRaisesRegex(InputError, 'encrypted.*password'):
            load_sources([encrypted])

    @unittest.skipUnless(HAS_DOCX, 'python-docx is optional')
    def test_docx_preserves_paragraph_table_order_and_heading(self):
        from docx import Document
        path = self.root / 'ordered.docx'
        document = Document()
        document.core_properties.title = 'Synthetic Word test'
        document.add_heading('Finding', level=1)
        document.add_paragraph('Before table.')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Name'
        table.cell(0, 1).text = 'Value'
        table.cell(1, 0).text = 'Alpha'
        table.cell(1, 1).text = '3'
        document.add_paragraph('After table.')
        document.save(path)
        sources, provenance = load_sources([path])
        text = sources[0]['text']
        self.assertIn('# Finding', text)
        self.assertLess(text.index('Before table.'), text.index('Alpha'))
        self.assertLess(text.index('Alpha'), text.index('After table.'))
        self.assertIn('| Name | Value |', text)
        self.assertIn('| Alpha | 3 |', text)
        extraction = provenance['sources'][0]['extraction']
        self.assertEqual(extraction['table_count'], 1)
        self.assertEqual(extraction['metadata']['title'], 'Synthetic Word test')

    @unittest.skipUnless(HAS_DOCX, 'python-docx is optional')
    def test_docx_empty_tables_do_not_count_as_extracted_text(self):
        from docx import Document
        path = self.root / 'empty-table.docx'
        document = Document()
        document.add_table(rows=1, cols=1)
        document.save(path)
        with self.assertRaisesRegex(InputError, 'no extractable main-document text'):
            load_sources([path])

    @unittest.skipUnless(HAS_DOCX, 'python-docx is optional')
    def test_docx_office_math_requires_editable_reconstruction(self):
        from docx import Document
        from docx.oxml import parse_xml
        path = self.root / 'math.docx'
        document = Document()
        paragraph = document.add_paragraph('Equation: ')
        paragraph._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x+1</m:t></m:r></m:oMath>'))
        document.save(path)
        sources, provenance = load_sources([path])
        self.assertIn('Equation: [待恢复为可编辑公式：S001-M001]', sources[0]['text'])
        self.assertNotIn('Equation: x+1', sources[0]['text'])
        self.assertEqual(provenance['sources'][0]['extraction']['math_count'], 1)
        finding = next(item for item in provenance['preparation_findings'] if item['code'] == 'docx_omml_requires_editable_math')
        self.assertEqual(finding['linear_text'], 'x+1')
        self.assertIn('<m:oMath', finding['original_omml'])
        self.assertEqual(finding['status'], 'unresolved')
        self.assertFalse(provenance['preparation']['equation_table_screenshots_allowed'])

    @unittest.skipUnless(HAS_DOCX, 'python-docx is optional')
    def test_docx_original_image_resources_materialize_with_cjk_text(self):
        from docx import Document
        document = Document()
        document.add_paragraph('中文图注：原始插图。')
        document.add_picture(io.BytesIO(tiny_png()))
        path = self.root / '中文原稿.docx'
        document.save(path)
        prepared = self.root / '准备目录'
        sources, provenance = load_sources([path], prepared_dir=prepared)
        self.assertIn('中文图注', sources[0]['text'])
        resource = provenance['resources'][0]
        self.assertEqual(resource['status'], 'materialized')
        self.assertTrue(resource['output_relative_path'].startswith('resources/S001/docx-image/'))
        self.assertIn('](' + resource['output_relative_path'] + ')', sources[0]['text'])
        self.assertEqual(Path(resource['path']).read_bytes(), tiny_png())
        record = next(item for item in provenance['files'] if item['role'] == 'docx-image')
        self.assertEqual(record['output_relative_path'], resource['output_relative_path'])
        self.assertEqual(record['sha256'], hashlib.sha256(tiny_png()).hexdigest())
        self.assertTrue(resource['requires_visual_classification'])

    @unittest.skipUnless(HAS_DOCX, 'python-docx is optional')
    def test_default_docx_image_read_is_read_only_and_records_pending_materialization(self):
        from docx import Document
        document = Document()
        document.add_paragraph('图片资源。')
        document.add_picture(io.BytesIO(tiny_png()))
        path = self.root / 'image.docx'
        document.save(path)
        before = set(self.root.iterdir())
        _, provenance = load_sources([path])
        self.assertEqual(set(self.root.iterdir()), before)
        self.assertEqual(provenance['resources'][0]['status'], 'needs_materialization')
        self.assertFalse(any(item['role'] == 'docx-image' for item in provenance['files']))
        self.assertTrue(any(item['code'] == 'image_resource_not_materialized' for item in provenance['preparation_findings']))

    @unittest.skipUnless(HAS_DOCX, 'python-docx is optional')
    def test_docx_merged_table_requires_structural_resolution(self):
        from docx import Document
        document = Document()
        document.add_paragraph('合并表格。')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = '合并标题'
        table.cell(1, 0).text = '甲'
        table.cell(1, 1).text = '乙'
        path = self.root / 'merged.docx'
        document.save(path)
        sources, provenance = load_sources([path])
        self.assertIn('合并标题', sources[0]['text'])
        finding = next(item for item in provenance['preparation_findings'] if item['code'] == 'docx_table_requires_review')
        self.assertTrue(finding['complex_structure'])
        self.assertEqual(finding['status'], 'unresolved')
        self.assertIn('gridSpan', finding['original_table_xml'])

    @unittest.skipUnless(HAS_DOCX, 'python-docx is optional')
    def test_explicit_pandoc_math_path_still_requires_visual_review(self):
        from docx import Document
        from docx.oxml import parse_xml
        document = Document()
        paragraph = document.add_paragraph('中文公式：')
        paragraph._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:r><m:t>x</m:t></m:r></m:num><m:den><m:r><m:t>y</m:t></m:r></m:den></m:f></m:oMath>'))
        document.add_picture(io.BytesIO(tiny_png()))
        path = self.root / 'math-rich.docx'
        document.save(path)
        result = subprocess.CompletedProcess([], 0,
            stdout='中文公式：$\\frac{x}{y}$\n\n![](media/image1.png)\n'.encode('utf-8'), stderr=b'')
        with patch('readeraware.inputs.shutil.which', return_value='C:/fake/pandoc.exe'), patch('readeraware.inputs.subprocess.run', return_value=result) as conversion:
            sources, provenance = load_sources([path], docx_engine='pandoc', prepared_dir=self.root / 'prepared')
        self.assertIn('$\\frac{x}{y}$', sources[0]['text'])
        self.assertIn('](resources/S001/docx-image/image-0001.png)', sources[0]['text'])
        finding = next(item for item in provenance['preparation_findings'] if item['code'] == 'docx_omml_requires_editable_math')
        self.assertEqual(finding['status'], 'needs_visual_verification')
        self.assertFalse(provenance['preparation']['editable_markdown_verified'])
        command = conversion.call_args.args[0]
        self.assertNotIn('--extract-media', ' '.join(command))
        self.assertFalse(conversion.call_args.kwargs['shell'])

    @unittest.skipUnless(HAS_DOCX, 'python-docx is optional')
    def test_explicit_pandoc_missing_dependency_fails_clearly(self):
        from docx import Document
        document = Document()
        document.add_paragraph('中文正文。')
        path = self.root / 'source.docx'
        document.save(path)
        with patch('readeraware.inputs.shutil.which', return_value=None):
            with self.assertRaisesRegex(InputError, '未找到本地 pandoc'):
                load_sources([path], docx_engine='pandoc', pandoc_bin='missing-explicit-pandoc')

    @unittest.skipUnless(HAS_DOCX, 'python-docx is optional')
    def test_pandoc_missing_math_and_unmatched_images_remain_unresolved(self):
        from docx import Document
        from docx.oxml import parse_xml
        document = Document()
        paragraph = document.add_paragraph('待转换公式。')
        paragraph._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x+1</m:t></m:r></m:oMath>'))
        path = self.root / 'incomplete.docx'
        document.save(path)
        result = subprocess.CompletedProcess([], 0,
            stdout='中文正文。 ![](https://example.invalid/missing.png)'.encode('utf-8'), stderr=b'')
        with patch('readeraware.inputs.shutil.which', return_value='C:/fake/pandoc.exe'), patch('readeraware.inputs.subprocess.run', return_value=result):
            _, provenance = load_sources([path], docx_engine='pandoc')
        unresolved = {finding['code'] for finding in provenance['preparation_findings'] if finding['status'] == 'unresolved'}
        self.assertIn('pandoc_math_conversion_incomplete', unresolved)
        self.assertIn('pandoc_image_reference_unresolved', unresolved)
        self.assertIn('docx_omml_requires_editable_math', unresolved)
        self.assertFalse(provenance['preparation']['editable_markdown_verified'])

    @unittest.skipUnless(HAS_DOCX, 'python-docx is optional')
    def test_resource_collision_is_not_silently_overwritten(self):
        from docx import Document
        document = Document()
        document.add_paragraph('资源冲突测试。')
        document.add_picture(io.BytesIO(tiny_png()))
        path = self.root / 'source.docx'
        document.save(path)
        prepared = self.root / 'prepared'
        target = prepared / 'resources/S001/docx-image/image-0001.png'
        target.parent.mkdir(parents=True)
        target.write_bytes(b'other bytes')
        with self.assertRaisesRegex(InputError, '内容不同'):
            load_sources([path], prepared_dir=prepared)
        self.assertEqual(target.read_bytes(), b'other bytes')

    @unittest.skipUnless(HAS_PDF and HAS_REPORTLAB, 'pypdf and reportlab are optional')
    def test_pdf_embedded_images_are_local_resources_not_fidelity_claims(self):
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.lib.utils import ImageReader
        path = self.root / 'image.pdf'
        canvas = Canvas(str(path))
        canvas.drawString(40, 750, 'Synthetic image and editable text.')
        canvas.drawImage(ImageReader(io.BytesIO(tiny_png())), 40, 600, width=80, height=80)
        canvas.save()
        sources, provenance = load_sources([path], prepared_dir=self.root / 'prepared')
        resource = provenance['resources'][0]
        self.assertEqual(resource['kind'], 'pdf-image')
        self.assertEqual(resource['page_number'], 1)
        self.assertIn('](' + resource['output_relative_path'] + ')', sources[0]['text'])
        self.assertTrue(Path(resource['path']).is_file())
        self.assertTrue(any(item['code'] == 'pdf_editable_fidelity_unknown' and item['status'] == 'unresolved' for item in provenance['preparation_findings']))


if __name__ == '__main__':
    unittest.main()
